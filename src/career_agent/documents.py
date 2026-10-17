from __future__ import annotations

from io import BytesIO
import re
import zipfile
from typing import Any

from .models import ApplicationPack, Job
from .profile import CandidateProfile


def build_cv_docx(pack: ApplicationPack, job: Job, profile: CandidateProfile) -> bytes:
    if pack.rich_content:
        from .cv_layout import render_rich_cv
        return render_rich_cv(pack, profile)
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.right_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    _set_style(styles["Normal"], "Arial", 10.2, "1F2430", after=3, line_spacing=1.05)
    _set_style(styles["Heading 1"], "Arial", 11.5, "372F7A", bold=True, before=9, after=3)
    _set_style(styles["Heading 2"], "Arial", 10.5, "11131A", bold=True, before=4, after=1)
    _set_style(styles["List Bullet"], "Arial", 9.6, "1F2430", after=1, line_spacing=1.0)

    identity = profile.data["identity"]
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run(identity["name"])
    _set_run(run, "Arial", 21, "11131A", bold=True)

    headline = doc.add_paragraph()
    headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    headline.paragraph_format.space_after = Pt(3)
    run = headline.add_run(pack.cv_headline)
    _set_run(run, "Arial", 10.5, "372F7A", bold=True)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(6)
    contact_text = " | ".join(
        [identity["location"], identity["email"], identity["linkedin"], identity["github"]]
    )
    _set_run(contact.add_run(contact_text), "Arial", 8.5, "555B66")

    _add_heading(doc, "PROFILE")
    doc.add_paragraph(pack.cv_summary)

    evidence_by_id = {item["id"]: item for item in profile.evidence}
    selected = set(pack.selected_evidence_ids)
    _add_heading(doc, "EXPERIENCE")
    for role in profile.data["experience"]:
        heading = doc.add_paragraph(style="Heading 2")
        _set_run(heading.add_run(f"{role['title']} — {role['company']}"), "Arial", 10.5, "11131A", bold=True)
        dates = heading.add_run(f"   {role['from']} – {role['to']} | {role['location']}")
        _set_run(dates, "Arial", 8.3, "676D78", italic=True)
        role_evidence = [evidence_by_id[item] for item in role["evidence_ids"] if item in evidence_by_id]
        role_evidence.sort(key=lambda item: item["id"] not in selected)
        for item in role_evidence:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.first_line_indent = Inches(-0.12)
            paragraph.add_run(_tailored_claim(item["claim"], pack.tailored_bullets))

    _add_heading(doc, "SELECTED AI ENGINEERING PROJECTS")
    for evidence_id in ("EXAMPLE-APP", "EXAMPLE-MODEL", "EXAMPLE-TOOLS"):
        item = evidence_by_id[evidence_id]
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
        paragraph.add_run(item["claim"])

    _add_heading(doc, "TECHNICAL SKILLS")
    # Job skills describe what the employer wants, not what the candidate has.
    # Use them only to prioritise verified profile skills and never copy a
    # requirement into the CV as an unsupported capability.
    candidate_skills = list(
        dict.fromkeys(
            profile.data["skills"]["expert"]
            + profile.data["skills"]["strong"]
            + profile.data["skills"]["working"]
        )
    )
    requested = {skill.casefold() for skill in job.skills}
    selected_skills = sorted(
        candidate_skills,
        key=lambda skill: (skill.casefold() not in requested, candidate_skills.index(skill)),
    )
    doc.add_paragraph(" • ".join(selected_skills[:24]))

    _add_heading(doc, "EDUCATION")
    for education in profile.data["education"]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        _set_run(
            paragraph.add_run(f"{education['degree']} — {education['institution']}"),
            "Arial",
            9.5,
            "11131A",
            bold=True,
        )
        _set_run(
            paragraph.add_run(f" | {education['result']} | {education['from']}–{education['to']}"),
            "Arial",
            8.7,
            "555B66",
        )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(footer.add_run(f"Tailored for {job.title} at {job.company}"), "Arial", 7.5, "7B808A")

    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _build_rich_cv_docx(pack: ApplicationPack, profile: CandidateProfile) -> bytes:
    """Compatibility wrapper for callers of the earlier private renderer."""
    from .cv_layout import render_rich_cv
    return render_rich_cv(pack, profile)


def build_cover_letter_docx(pack: ApplicationPack, job: Job, profile: CandidateProfile) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.9)
    _set_style(doc.styles["Normal"], "Arial", 10.8, "1F2430", after=8, line_spacing=1.12)

    identity = profile.data["identity"]
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    _set_run(title.add_run(identity["name"]), "Arial", 19, "11131A", bold=True)
    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(22)
    _set_run(metadata.add_run(f"{identity['email']} | {identity['location']} | {identity['linkedin']}"), "Arial", 8.8, "666C78")

    for block in pack.cover_letter.split("\n\n"):
        paragraph = doc.add_paragraph(block.strip())
        paragraph.paragraph_format.space_after = Pt(10)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(footer.add_run(f"Application: {job.title} — {job.company}"), "Arial", 7.5, "7B808A")
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def scrub_docx_metadata(content: bytes) -> bytes:
    """Remove author metadata, custom properties, and Word revision session ids."""
    source = BytesIO(content)
    destination = BytesIO()
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for info in source_zip.infolist():
            if info.filename == "docProps/custom.xml":
                continue
            payload = source_zip.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = payload.decode("utf-8")
                text = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator></dc:creator>", text)
                text = re.sub(
                    r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>",
                    "<cp:lastModifiedBy></cp:lastModifiedBy>",
                    text,
                )
                payload = text.encode("utf-8")
            if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                payload = re.sub(rb'\s+w:rsid[A-Za-z]+="[0-9A-Fa-f]+"', b"", payload)
            target_zip.writestr(info, payload)
    return destination.getvalue()


def _set_style(
    style: Any,
    font: str,
    size: float,
    color: str,
    *,
    bold: bool = False,
    before: float = 0,
    after: float = 0,
    line_spacing: float = 1.0,
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    style.font.name = font
    style._element.rPr.rFonts.set(qn("w:ascii"), font)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line_spacing


def _set_run(
    run: Any,
    font: str,
    size: float,
    color: str,
    *,
    bold: bool = False,
    italic: bool = False,
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def _add_heading(doc: Any, text: str) -> None:
    paragraph = doc.add_paragraph(text, style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True


def _tailored_claim(original: str, tailored: list[str]) -> str:
    original_numbers = set(part for part in original.split() if any(char.isdigit() for char in part))
    for candidate in tailored:
        if original_numbers and original_numbers <= set(candidate.split()):
            return candidate
    return original
