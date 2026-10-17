"""Renderer for the owner's reviewed CV design; no inference or factual synthesis."""
from io import BytesIO
import re

from .application_writing import WORK


def skill_groups(skills):
    """Grouping changes labels/order only; never copies skills from a vacancy."""
    groups = {k: [] for k in ('Programming and data', 'Machine learning and AI', 'Cloud', 'Engineering tools', 'Other')}
    for skill in skills:
        text = skill['text']
        lower = text.casefold()
        if any(x in lower for x in ('aws', 'gcp', 'azure', 'bigquery', 'vertex', 'bedrock', 'textract', 'lambda')) and not lower.startswith('sql:'):
            group = 'Cloud'
        elif any(x in lower for x in ('python', 'sql', 'pandas', 'numpy', 'scipy')) or lower in ('r', 'c/c++', 'c#'):
            group = 'Programming and data'
        elif any(x in lower for x in ('docker', 'git', 'terraform', 'dbt', 'airbyte', 'hadoop', 'grafana', 'power bi', 'matplotlib', 'seaborn')):
            group = 'Engineering tools'
        elif lower in ('english', 'french', 'persian'):
            group = 'Other'
        else:
            group = 'Machine learning and AI'
        groups[group].append(text)
    return {k: list(dict.fromkeys(v)) for k, v in groups.items() if v}


def display_date(value):
    if re.fullmatch(r'\d{4}-\d{2}', value or ''):
        from datetime import datetime
        return datetime.strptime(value, '%Y-%m').strftime('%b %Y')
    return value or ''


def render_rich_cv(pack, profile):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(17)
    section.left_margin = section.right_margin = Mm(18)
    blue = '326CA2'
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for name, size in [('Normal', 11), ('Title', 28), ('Heading 1', 14), ('Heading 2', 11.5), ('List Bullet', 11)]:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.07
    for name in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.keep_with_next = True
    for name in ('Title', 'Heading 1'):
        doc.styles[name].font.color.rgb = RGBColor.from_string(blue)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
    doc.styles['List Bullet'].paragraph_format.left_indent = Mm(5)
    doc.styles['List Bullet'].paragraph_format.first_line_indent = Mm(-3.5)
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), '42')
    level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
    for name, value in [('start', '1'), ('numFmt', 'bullet'), ('lvlText', '•'), ('lvlJc', 'left'), ('suff', 'tab')]:
        el = OxmlElement('w:' + name); el.set(qn('w:val'), value); level.append(el)
    props = OxmlElement('w:pPr'); tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'num'); tab.set(qn('w:pos'), '284'); tabs.append(tab); props.append(tabs)
    indent = OxmlElement('w:ind'); indent.set(qn('w:left'), '284'); indent.set(qn('w:hanging'), '199'); props.append(indent)
    level.append(props); rp = OxmlElement('w:rPr'); fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Arial'); fonts.set(qn('w:hAnsi'), 'Arial'); rp.append(fonts)
    color = OxmlElement('w:color'); color.set(qn('w:val'), blue); rp.append(color); level.append(rp)
    abstract.append(level); numbering.append(abstract)
    num = OxmlElement('w:num'); num.set(qn('w:numId'), '42')
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), '42'); num.append(ref); numbering.append(num)

    def paragraph(text='', style=None):
        p = doc.add_paragraph(text, style)
        p.paragraph_format.keep_together = True
        return p

    def heading(text):
        p = paragraph(text, 'Heading 1')
        borders = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
        for k, v in [('val', 'single'), ('sz', '4'), ('space', '4'), ('color', 'A4B8CA')]:
            bottom.set(qn('w:' + k), v)
        borders.append(bottom); p._p.get_or_add_pPr().append(borders)
        return p

    def record_heading(title, dates='', subtitle=''):
        p = paragraph('', 'Heading 2')
        # Avoid long-title/date collision; fall back to a separate metadata line.
        if len(title) <= 48 and len(dates) <= 27:
            p.paragraph_format.tab_stops.add_tab_stop(Mm(174), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(title).bold = True
            r = p.add_run('\t' + dates); r.bold = False; r.font.size = Pt(10.5)
        else:
            p.add_run(title).bold = True
            subtitle = ' | '.join(x for x in (dates, subtitle) if x)
        if subtitle:
            p = paragraph(subtitle)
            p.runs[0].italic = True
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(3)

    identity = profile.data['identity']
    for text, style in [(profile.name, 'Title'), (pack.cv_headline, None),
        (' | '.join(identity[k] for k in ('location', 'email') if identity.get(k)), None),
        (' | '.join(identity[k] for k in ('linkedin', 'github') if identity.get(k)), None)]:
        p = paragraph(text, style); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        if not style: p.runs[0].font.size = Pt(10.5)
    for award in pack.rich_content.get('awards', []):
        p = paragraph(award['text']); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True; p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(9)
    if pack.cv_summary:
        paragraph(pack.cv_summary)
    groups = skill_groups(pack.rich_content.get('skills', []))
    if groups:
        heading('Technical Skills')
        for label, values in groups.items():
            p = paragraph(); p.add_run(label + ': ').bold = True
            p.add_run(', '.join(values))

    sequence = [('Professional Experience', True), ('Projects', False)]
    if pack.rich_content.get('section_order') == 'projects_first':
        sequence.reverse()
    for name, is_work in sequence:
        rows = [r for r in pack.rich_content['sections'] if (r['category'] in WORK) == is_work]
        if not rows: continue
        heading(name)
        for row in rows:
            dates = row.get('dates') or {}
            end = 'Present' if dates.get('ongoing') else display_date(dates.get('end'))
            date_text = ' - '.join(v for v in (display_date(dates.get('start')), end) if v)
            if is_work:
                company = (row.get('organization') or {}).get('value')
                role = (row.get('role_or_degree') or {}).get('value')
                location = (row.get('location') or {}).get('value')
                record_heading(company or row['title'], date_text, ' | '.join(v for v in (role, location) if v))
            else:
                scope = 'University project' if row['category'] == 'university_project' else 'Personal project'
                record_heading(row['title'], scope)
            for b in row['bullets']:
                p = paragraph('', 'List Bullet')
                np = OxmlElement('w:numPr'); lev = OxmlElement('w:ilvl'); lev.set(qn('w:val'), '0'); np.append(lev)
                ni = OxmlElement('w:numId'); ni.set(qn('w:val'), '42'); np.append(ni); p._p.get_or_add_pPr().append(np)
                # Emphasise existing numeric outcomes only, never add/rewrite their wording.
                for part in re.split(r'([$£€]?\b\d[\d,.]*\+?%?(?:\s+ms\b)?)', b['text']):
                    p.add_run(part).bold = bool(re.fullmatch(r'[$£€]?\d[\d,.]*\+?%?(?:\s+ms)?', part))
    for name, category in [('Education', 'education'), ('Certifications', 'certification')]:
        rows = [r for r in pack.rich_content['qualifications'] if r['category'] == category]
        if not rows: continue
        heading(name)
        for row in rows:
            dates = row.get('dates') or {}
            if category == 'education':
                date_text = ' - '.join(display_date(dates[k]) for k in ('start', 'end') if dates.get(k))
                record_heading(row['organization']['value'], date_text, row['role_or_degree']['value'])
                doc.paragraphs[-1].paragraph_format.keep_with_next = False
                doc.paragraphs[-1].paragraph_format.space_after = Pt(10)
            else:
                date_text = display_date(dates.get('start'))
                if dates.get('end'): date_text += '; valid until ' + display_date(dates['end'])
                paragraph(row['title'] + (' | ' + date_text if date_text else ''))
    for key in ('author', 'last_modified_by', 'comments', 'keywords', 'subject'):
        setattr(doc.core_properties, key, '')
    doc.core_properties.title = profile.name + ' CV'
    output = BytesIO(); doc.save(output)
    return output.getvalue()
